import os
import re
import base64
from fuzzywuzzy import fuzz
import openai

def sanitize_filename(filename):
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', '_', filename)

def extract_file_paths(text):
    return re.findall(r'Fig_path=\s*"(.*?)"', text)

def encode_image_to_base64(image_path):
    try:
        with open(image_path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode("utf-8")
    except Exception as e:
        print(f"Failed to encode {image_path}: {e}")
        return None

def extract_relevant_sections(full_text, prompt, threshold=70):
    sections = full_text.split("\n\n")
    return "\n\n".join([s for s in sections if fuzz.partial_ratio(prompt.lower(), s.lower()) > threshold])

def classify_pdf(path, threshold=0.1):
    import fitz
    doc = fitz.open(path)
    text_pages = sum(1 for p in doc if sum(len(b[4].strip()) for b in p.get_text("blocks")) > threshold)
    doc.close()
    return "Text-based PDF" if text_pages > len(doc) // 2 else "Image-based PDF"

def extract_images_from_docx(path):
    from docx import Document
    doc = Document(path)
    images = {}
    text = []
    fig_pat = re.compile(r"\(Fig\.\d+\)")
    folder = os.path.join(os.getcwd(), os.path.splitext(os.path.basename(path))[0], "images")
    os.makedirs(folder, exist_ok=True)

    new_doc = Document()
    img_counter = 1

    for para in doc.paragraphs:
        for run in para.runs:
            if run._element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"):
                for shape in run._element.findall(".//{http://schemas.openxmlformats.org/drawing/2006/main}blip"):
                    embed_id = shape.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"]
                    part = doc.part.related_parts[embed_id]
                    filename = f"extracted_image{img_counter}_Fig.{img_counter}.png"
                    img_path = os.path.join(folder, filename)
                    with open(img_path, "wb") as f:
                        f.write(part.blob)
                    rel_path = img_path.replace("\\", "/")
                    fig_ref = f'Refer Fig_path= "{rel_path}"'
                    para.text = fig_pat.sub(fig_ref, para.text) if fig_pat.search(para.text) else para.text + "\n" + fig_ref
                    images[f"Image_{img_counter}"] = rel_path
                    img_counter += 1
        new_doc.add_paragraph(para.text)
    new_doc_path = os.path.join(os.path.dirname(folder), os.path.basename(path).replace(".docx", "_modified.docx"))
    new_doc.save(new_doc_path)
    return images, "\n".join([p.text for p in new_doc.paragraphs])

def extract_images_from_pdf(path, format="txt"):
    import fitz
    from docx import Document

    doc = fitz.open(path)
    folder = os.path.join(os.getcwd(), os.path.splitext(os.path.basename(path))[0], "images")
    os.makedirs(folder, exist_ok=True)

    images = {}
    text_out = []
    img_counter = 1
    fig_pat = re.compile(r"\(Fig\.\d+\)")

    for page in doc:
        lines = page.get_text("text").split("\n")
        for img in page.get_images(full=True):
            xref = img[0]
            base = doc.extract_image(xref)
            filename = f"extracted_image{img_counter}_Fig.{img_counter}.{base['ext']}"
            path_img = os.path.join(folder, filename)
            with open(path_img, "wb") as f:
                f.write(base["image"])
            rel_path = path_img.replace("\\", "/")
            ref = f'Refer Fig_path= "{rel_path}"'
            for i, line in enumerate(lines):
                if fig_pat.search(line):
                    lines[i] = fig_pat.sub(ref, line)
                    break
            else:
                lines.append(ref)
            images[f"Image_{img_counter}"] = rel_path
            img_counter += 1
        text_out.append("\n".join(lines))
    doc.close()
    return images, "\n".join(text_out)

def extract_images_from_scan_pdf(path, format="txt"):
    from pdf2image import convert_from_path
    import numpy as np
    import pytesseract
    from PIL import Image
    import cv2

    filename = os.path.splitext(os.path.basename(path))[0]
    folder = os.path.join(os.getcwd(), filename, "images")
    os.makedirs(folder, exist_ok=True)

    images = {}
    output_text = []
    img_counter = 1
    fig_pat = re.compile(r"\(Fig\.\d+\)")

    for page in convert_from_path(path, dpi=300):
        img_cv = cv2.cvtColor(np.array(page), cv2.COLOR_RGB2BGR)
        text = pytesseract.image_to_string(img_cv)
        lines = text.split("\n")
        gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
        _, thresh = cv2.threshold(gray, 200, 255, cv2.THRESH_BINARY_INV)
        contours, _ = cv2.findContours(thresh, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

        for cnt in contours:
            x, y, w, h = cv2.boundingRect(cnt)
            if w * h > 5000:
                crop = img_cv[y:y+h, x:x+w]
                pil_img = Image.fromarray(cv2.cvtColor(crop, cv2.COLOR_BGR2RGB))
                filename = f"extracted_image{img_counter}_Fig.{img_counter}.png"
                path_img = os.path.join(folder, filename)
                pil_img.save(path_img)
                rel_path = path_img.replace("\\", "/")
                ref = f'Refer Fig_path= "{rel_path}"'
                for i, line in enumerate(lines):
                    if fig_pat.search(line):
                        lines[i] = fig_pat.sub(ref, line)
                        break
                else:
                    lines.append(ref)
                images[f"Image_{img_counter}"] = rel_path
                img_counter += 1
        output_text.append("\n".join(lines))
    return images, "\n".join(output_text)

def get_openai_client():
    from dotenv import load_dotenv
    load_dotenv()
    return openai.AzureOpenAI(
        api_key=os.getenv("OPENAI_API_KEY"),
        api_version=os.getenv("OPENAI_API_VERSION"),
        azure_endpoint=os.getenv("AZURE_OPENAI_ENDPOINT")
    )
